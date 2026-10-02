"""
Nóminas: cálculo mensual, recibos en PDF, remesa SEPA de pago y asiento
contable.

Es un cálculo de gestión para preparar y revisar nóminas estándar (salario
fijo, pagas extra, horas extra e incentivos). No cubre incapacidad
temporal, embargos, atrasos ni particularidades de cada convenio: la nómina
oficial debe validarla la persona responsable o la gestoría.
"""
from __future__ import annotations

from app import clock
import calendar
import io
import re
from dataclasses import dataclass
from datetime import date
from datetime import datetime
from datetime import timezone
from decimal import Decimal
from decimal import ROUND_HALF_UP
from typing import Any
from xml.sax.saxutils import escape

from sqlalchemy import select
from sqlalchemy.orm import Session
from sqlalchemy.orm import selectinload

from app.models import CompanyProfile
from app.models import Employee
from app.models import PayrollRun
from app.models import Payslip


CENT = Decimal("0.01")
ZERO = Decimal("0.00")

MONTH_NAMES = (
    "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
    "agosto", "septiembre", "octubre", "noviembre", "diciembre",
)

EXTRA_PAY_MONTHS = {6, 12}

TEMPORARY_CONTRACTS = {"TEMPORAL", "PRACTICAS", "FORMACION"}


@dataclass(frozen=True)
class ContributionParams:
    # Porcentajes a cargo del trabajador.
    employee_common: Decimal
    employee_unemployment_permanent: Decimal
    employee_unemployment_temporary: Decimal
    employee_training: Decimal
    employee_mei: Decimal
    # Porcentajes a cargo de la empresa.
    employer_common: Decimal
    employer_unemployment_permanent: Decimal
    employer_unemployment_temporary: Decimal
    employer_fogasa: Decimal
    employer_training: Decimal
    employer_mei: Decimal
    # Bases mensuales (grupo general).
    base_min: Decimal
    base_max: Decimal
    source: str


# Tipos generales del Régimen General. El MEI sube cada año (Ley 21/2021).
# Las bases mínima y máxima se revisan cada año por orden ministerial:
# compruébalas y actualiza esta tabla al publicarse.
PARAMS: dict[int, ContributionParams] = {
    2025: ContributionParams(
        employee_common=Decimal("4.70"),
        employee_unemployment_permanent=Decimal("1.55"),
        employee_unemployment_temporary=Decimal("1.60"),
        employee_training=Decimal("0.10"),
        employee_mei=Decimal("0.13"),
        employer_common=Decimal("23.60"),
        employer_unemployment_permanent=Decimal("5.50"),
        employer_unemployment_temporary=Decimal("6.70"),
        employer_fogasa=Decimal("0.20"),
        employer_training=Decimal("0.60"),
        employer_mei=Decimal("0.67"),
        base_min=Decimal("1381.20"),
        base_max=Decimal("4909.50"),
        source="Tipos 2025 · bases mínima y máxima 2025",
    ),
    2026: ContributionParams(
        employee_common=Decimal("4.70"),
        employee_unemployment_permanent=Decimal("1.55"),
        employee_unemployment_temporary=Decimal("1.60"),
        employee_training=Decimal("0.10"),
        employee_mei=Decimal("0.15"),
        employer_common=Decimal("23.60"),
        employer_unemployment_permanent=Decimal("5.50"),
        employer_unemployment_temporary=Decimal("6.70"),
        employer_fogasa=Decimal("0.20"),
        employer_training=Decimal("0.60"),
        employer_mei=Decimal("0.75"),
        base_min=Decimal("1381.20"),
        base_max=Decimal("4909.50"),
        source=(
            "Tipos 2026 (MEI 0,90 %) · bases de 2025 pendientes de "
            "actualizar con la orden de cotización de 2026"
        ),
    ),
}

DEFAULT_AT_EP = Decimal("1.50")


def params_for(year: int) -> ContributionParams:
    if year in PARAMS:
        return PARAMS[year]

    return PARAMS[max(PARAMS)] if year > max(PARAMS) else PARAMS[min(PARAMS)]


def money(value: Decimal | float | int | None) -> Decimal:
    if value is None:
        return ZERO

    return Decimal(str(value)).quantize(CENT, rounding=ROUND_HALF_UP)


def pct(base: Decimal, rate: Decimal) -> Decimal:
    return money(base * rate / Decimal(100))


# -------------------------------------------------------------------
# IRPF (estimación)
# -------------------------------------------------------------------

IRPF_SCALE = (
    (Decimal("12450"), Decimal("19")),
    (Decimal("20200"), Decimal("24")),
    (Decimal("35200"), Decimal("30")),
    (Decimal("60000"), Decimal("37")),
    (Decimal("300000"), Decimal("45")),
    (None, Decimal("47")),
)

PERSONAL_MINIMUM = Decimal("5550")
CHILD_MINIMUMS = (Decimal("2400"), Decimal("2700"), Decimal("4000"), Decimal("4500"))
OTHER_EXPENSES = Decimal("2000")
NO_WITHHOLDING_LIMIT = Decimal("15876")


def apply_scale(base: Decimal) -> Decimal:
    tax = ZERO
    lower = ZERO

    for upper, rate in IRPF_SCALE:
        if base <= lower:
            break

        top = base if upper is None else min(base, upper)
        tax += (top - lower) * rate / Decimal(100)

        if upper is None or base <= upper:
            break

        lower = upper

    return tax


def work_income_reduction(net: Decimal) -> Decimal:
    # Art. 20 LIRPF (redacción vigente desde 2024).
    if net <= Decimal("14852"):
        return Decimal("7302")

    if net <= Decimal("17673.52"):
        return Decimal("7302") - Decimal("1.75") * (net - Decimal("14852"))

    if net <= Decimal("19747.5"):
        return Decimal("2364.34") - Decimal("1.14") * (net - Decimal("17673.52"))

    return ZERO


def estimate_irpf_rate(
    *,
    annual_gross: Decimal,
    annual_ss: Decimal,
    children: int = 0,
    temporary: bool = False,
) -> Decimal:
    """
    Tipo de retención aproximado (escala general estatal + autonómica
    tipo). La empresa debe aplicar el cálculo oficial de la AEAT.
    """
    if annual_gross <= 0:
        return ZERO

    minimum = PERSONAL_MINIMUM + sum(CHILD_MINIMUMS[: min(children, 4)], ZERO)
    minimum += CHILD_MINIMUMS[-1] * max(0, children - 4)

    if annual_gross <= NO_WITHHOLDING_LIMIT and children == 0:
        rate = ZERO
    else:
        net = annual_gross - annual_ss
        base = max(ZERO, net - OTHER_EXPENSES - work_income_reduction(net))
        tax = max(ZERO, apply_scale(base) - apply_scale(min(minimum, base)))
        rate = (tax / annual_gross * Decimal(100)).quantize(CENT)

    if temporary:
        rate = max(rate, Decimal("2.00"))

    return rate


# -------------------------------------------------------------------
# Cálculo de un recibo
# -------------------------------------------------------------------

def employed_days(employee: Employee, year: int, month: int) -> int:
    """Días trabajados en el mes, sobre 30 (mes comercial)."""
    start = date(year, month, 1)
    end = date(year, month, calendar.monthrange(year, month)[1])

    first = max(start, employee.hire_date) if employee.hire_date else start
    last = min(end, employee.termination_date) if employee.termination_date else end

    if first > last:
        return 0

    if first == start and last == end:
        return 30

    return min(30, (last - first).days + 1)


def calculate_payslip(
    employee: Employee,
    *,
    year: int,
    month: int,
    overtime: Decimal = ZERO,
    bonus: Decimal = ZERO,
    advance: Decimal = ZERO,
    at_ep_rate: Decimal = DEFAULT_AT_EP,
) -> dict[str, Any]:
    params = params_for(year)
    days = employed_days(employee, year, month)
    factor = Decimal(days) / Decimal(30)

    annual = money(employee.annual_salary)
    payments = 14 if employee.payments_per_year == 14 else 12
    monthly = annual / Decimal(payments)
    temporary = employee.contract_type in TEMPORARY_CONTRACTS

    earnings: list[dict[str, Any]] = []
    base_salary = money(monthly * factor)
    earnings.append(
        {
            "concept": "Salario base" + (f" ({days}/30 días)" if days < 30 else ""),
            "amount": float(base_salary),
        }
    )

    extra = ZERO

    if payments == 14 and month in EXTRA_PAY_MONTHS:
        extra = money(monthly * factor)
        earnings.append(
            {
                "concept": "Paga extraordinaria de " + ("verano" if month == 6 else "Navidad"),
                "amount": float(extra),
            }
        )

    overtime = money(overtime)
    bonus = money(bonus)

    if overtime:
        earnings.append({"concept": "Horas extraordinarias", "amount": float(overtime)})

    if bonus:
        earnings.append({"concept": "Incentivos / complementos", "amount": float(bonus)})

    gross = base_salary + extra + overtime + bonus

    # Base de cotización: salario mensual con pagas extra prorrateadas.
    prorated = annual / Decimal(12) * factor
    base = money(prorated + overtime + bonus)
    base_min = money(params.base_min * Decimal(employee.workday_percent) / Decimal(100) * factor)
    base_max = money(params.base_max * factor)
    capped_base = min(max(base, base_min), base_max) if days else ZERO

    unemployment_employee = (
        params.employee_unemployment_temporary
        if temporary
        else params.employee_unemployment_permanent
    )
    unemployment_employer = (
        params.employer_unemployment_temporary
        if temporary
        else params.employer_unemployment_permanent
    )

    employee_lines = [
        ("Contingencias comunes", params.employee_common),
        ("Desempleo", unemployment_employee),
        ("Formación profesional", params.employee_training),
        ("Mecanismo de equidad intergeneracional (MEI)", params.employee_mei),
    ]
    employer_lines = [
        ("Contingencias comunes", params.employer_common),
        ("Desempleo", unemployment_employer),
        ("FOGASA", params.employer_fogasa),
        ("Formación profesional", params.employer_training),
        ("Accidentes de trabajo y enfermedad profesional", at_ep_rate),
        ("MEI", params.employer_mei),
    ]

    deductions = [
        {"concept": f"{name} ({rate} %)", "amount": float(pct(capped_base, rate)), "rate": float(rate)}
        for name, rate in employee_lines
    ]
    ss_employee = sum((pct(capped_base, rate) for _name, rate in employee_lines), ZERO)

    employer = [
        {"concept": f"{name} ({rate} %)", "amount": float(pct(capped_base, rate)), "rate": float(rate)}
        for name, rate in employer_lines
    ]
    ss_employer = sum((pct(capped_base, rate) for _name, rate in employer_lines), ZERO)

    if employee.irpf_rate is not None:
        irpf_rate = money(employee.irpf_rate)
        irpf_source = "manual"
    else:
        annual_ss = (
            min(max(annual / Decimal(12), params.base_min), params.base_max)
            * Decimal(12)
            * sum((rate for _name, rate in employee_lines), ZERO)
            / Decimal(100)
        )
        irpf_rate = estimate_irpf_rate(
            annual_gross=annual,
            annual_ss=annual_ss,
            children=employee.children or 0,
            temporary=temporary,
        )
        irpf_source = "estimado"

    irpf = pct(gross, irpf_rate)
    deductions.append(
        {
            "concept": f"Retención IRPF ({irpf_rate} %{' estimado' if irpf_source == 'estimado' else ''})",
            "amount": float(irpf),
            "rate": float(irpf_rate),
        }
    )

    advance = money(advance)

    if advance:
        deductions.append({"concept": "Anticipos", "amount": float(advance)})

    net = gross - ss_employee - irpf - advance

    return {
        "days": days,
        "gross": gross,
        "contribution_base": capped_base,
        "ss_employee": ss_employee,
        "irpf_rate": irpf_rate,
        "irpf_source": irpf_source,
        "irpf": irpf,
        "net": net,
        "ss_employer": ss_employer,
        "company_cost": gross + ss_employer,
        "lines": {
            "earnings": earnings,
            "deductions": deductions,
            "employer": employer,
            "base": {
                "computed": float(base),
                "applied": float(capped_base),
                "min": float(base_min),
                "max": float(base_max),
            },
            "params_source": params.source,
            "at_ep_rate": float(at_ep_rate),
        },
    }


def simulate_employee_cost(
    employee: Employee,
    *,
    year: int,
    at_ep_rate: Decimal = DEFAULT_AT_EP,
) -> dict[str, Any]:
    """Coste anual de un sueldo (simulador de contratación)."""
    totals = {key: ZERO for key in ("gross", "ss_employee", "irpf", "net", "ss_employer", "company_cost")}
    irpf_rate = ZERO

    for month in range(1, 13):
        slip = calculate_payslip(employee, year=year, month=month, at_ep_rate=at_ep_rate)
        irpf_rate = slip["irpf_rate"]

        for key in totals:
            totals[key] += slip[key]

    return {
        **{key: float(value) for key, value in totals.items()},
        "irpf_rate": float(irpf_rate),
        "monthly_net": float(money(totals["net"] / Decimal(employee.payments_per_year or 12))),
        "employer_ratio": (
            round(float(totals["ss_employer"] / totals["gross"] * 100), 1)
            if totals["gross"]
            else None
        ),
    }


# -------------------------------------------------------------------
# Procesos de nómina
# -------------------------------------------------------------------

def active_employees(database: Session, year: int, month: int) -> list[Employee]:
    employees = database.scalars(
        select(Employee).order_by(Employee.last_name, Employee.first_name)
    ).all()

    return [
        employee
        for employee in employees
        if employee.annual_salary and employed_days(employee, year, month) > 0
    ]


def employee_display_name(employee: Employee) -> str:
    return " ".join(part for part in (employee.first_name, employee.last_name) if part)


def company_at_ep(database: Session) -> Decimal:
    profile = database.scalar(select(CompanyProfile).limit(1))

    if profile and profile.at_ep_rate is not None:
        return Decimal(profile.at_ep_rate)

    return DEFAULT_AT_EP


def apply_slip(payslip: Payslip, employee: Employee, run: PayrollRun, at_ep: Decimal) -> None:
    result = calculate_payslip(
        employee,
        year=run.year,
        month=run.month,
        overtime=payslip.overtime or ZERO,
        bonus=payslip.bonus or ZERO,
        advance=payslip.advance or ZERO,
        at_ep_rate=at_ep,
    )
    payslip.employee_name = employee_display_name(employee)
    payslip.gross = result["gross"]
    payslip.contribution_base = result["contribution_base"]
    payslip.ss_employee = result["ss_employee"]
    payslip.irpf_rate = result["irpf_rate"]
    payslip.irpf = result["irpf"]
    payslip.net = result["net"]
    payslip.ss_employer = result["ss_employer"]
    payslip.company_cost = result["company_cost"]
    payslip.lines = {**result["lines"], "days": result["days"], "irpf_source": result["irpf_source"]}


def get_run(database: Session, run_id: int) -> PayrollRun | None:
    return database.scalar(
        select(PayrollRun)
        .where(PayrollRun.id == run_id)
        .options(selectinload(PayrollRun.payslips).selectinload(Payslip.employee))
    )


def create_run(database: Session, *, year: int, month: int) -> PayrollRun:
    existing = database.scalar(
        select(PayrollRun).where(PayrollRun.year == year, PayrollRun.month == month)
    )

    if existing is not None:
        raise ValueError("Ya existe la nómina de ese mes.")

    employees = active_employees(database, year, month)

    if not employees:
        raise ValueError(
            "No hay personas en plantilla con salario para ese mes. "
            "Añádelas en Equipo."
        )

    run = PayrollRun(year=year, month=month, status="DRAFT")
    database.add(run)
    database.flush()
    at_ep = company_at_ep(database)

    for employee in employees:
        payslip = Payslip(run=run, employee_id=employee.id, employee_name="")
        database.add(payslip)
        payslip.employee = employee
        apply_slip(payslip, employee, run, at_ep)

    database.flush()

    return run


def recalculate_run(database: Session, run: PayrollRun) -> PayrollRun:
    if run.status != "DRAFT":
        raise ValueError("Solo se puede recalcular una nómina en borrador.")

    at_ep = company_at_ep(database)
    current = {payslip.employee_id: payslip for payslip in run.payslips}

    for employee in active_employees(database, run.year, run.month):
        payslip = current.pop(employee.id, None)

        if payslip is None:
            payslip = Payslip(run=run, employee_id=employee.id, employee_name="")
            database.add(payslip)

        payslip.employee = employee
        apply_slip(payslip, employee, run, at_ep)

    for leftover in current.values():
        database.delete(leftover)

    database.flush()

    return run


def update_payslip_variables(
    database: Session,
    payslip: Payslip,
    changes: dict[str, Any],
) -> Payslip:
    if payslip.run.status != "DRAFT":
        raise ValueError("La nómina ya está aprobada: no se puede modificar.")

    for field in ("overtime", "bonus", "advance"):
        if field in changes and changes[field] is not None:
            value = money(changes[field])

            if value < 0:
                raise ValueError("Los importes no pueden ser negativos.")

            setattr(payslip, field, value)

    if payslip.employee is None:
        raise ValueError("La persona de este recibo ya no existe.")

    apply_slip(payslip, payslip.employee, payslip.run, company_at_ep(database))
    database.flush()

    return payslip


def run_totals(run: PayrollRun) -> dict[str, float]:
    keys = ("gross", "ss_employee", "irpf", "net", "ss_employer", "company_cost")

    return {
        key: float(sum((money(getattr(payslip, key)) for payslip in run.payslips), ZERO))
        for key in keys
    } | {"employees": len(run.payslips)}


def period_label(run: PayrollRun) -> str:
    return f"{MONTH_NAMES[run.month - 1]} {run.year}"


def serialize_payslip(payslip: Payslip) -> dict[str, Any]:
    employee = payslip.employee

    return {
        "id": payslip.id,
        "employee_id": payslip.employee_id,
        "employee_name": payslip.employee_name,
        "job_title": employee.job_title if employee else None,
        "iban_ok": bool(employee and employee.iban and iban_is_valid(employee.iban)),
        "overtime": float(payslip.overtime),
        "bonus": float(payslip.bonus),
        "advance": float(payslip.advance),
        "gross": float(payslip.gross),
        "contribution_base": float(payslip.contribution_base),
        "ss_employee": float(payslip.ss_employee),
        "irpf_rate": float(payslip.irpf_rate),
        "irpf": float(payslip.irpf),
        "net": float(payslip.net),
        "ss_employer": float(payslip.ss_employer),
        "company_cost": float(payslip.company_cost),
        "lines": payslip.lines,
    }


def serialize_run(run: PayrollRun, *, with_payslips: bool = False) -> dict[str, Any]:
    data: dict[str, Any] = {
        "id": run.id,
        "year": run.year,
        "month": run.month,
        "period": period_label(run),
        "status": run.status,
        "approved_at": run.approved_at.isoformat() if run.approved_at else None,
        "paid_at": run.paid_at.isoformat() if run.paid_at else None,
        "totals": run_totals(run),
        "journal": journal_entry(run),
    }

    if with_payslips:
        data["payslips"] = [serialize_payslip(payslip) for payslip in run.payslips]

    return data


def set_run_status(run: PayrollRun, status: str, paid_at: date | None = None) -> PayrollRun:
    transitions = {
        "APPROVED": {"DRAFT"},
        "PAID": {"APPROVED"},
        "DRAFT": {"APPROVED"},
    }

    if run.status not in transitions.get(status, set()):
        raise ValueError(f"No se puede pasar de {run.status} a {status}.")

    run.status = status

    if status == "APPROVED":
        run.approved_at = clock.now()
    elif status == "PAID":
        run.paid_at = paid_at or clock.today()
    elif status == "DRAFT":
        run.approved_at = None

    return run


# -------------------------------------------------------------------
# Asiento contable
# -------------------------------------------------------------------

def journal_entry(run: PayrollRun) -> list[dict[str, Any]]:
    totals = {key: Decimal(str(value)) for key, value in run_totals(run).items() if key != "employees"}
    advances = sum((money(payslip.advance) for payslip in run.payslips), ZERO)

    lines = [
        ("640", "Sueldos y salarios", totals["gross"], ZERO),
        ("642", "Seguridad Social a cargo de la empresa", totals["ss_employer"], ZERO),
        ("4751", "H.P. acreedora por retenciones practicadas", ZERO, totals["irpf"]),
        ("476", "Organismos de la Seguridad Social, acreedores", ZERO, totals["ss_employee"] + totals["ss_employer"]),
        ("465", "Remuneraciones pendientes de pago", ZERO, totals["net"]),
    ]

    if advances:
        lines.append(("460", "Anticipos de remuneraciones", ZERO, advances))

    return [
        {"account": account, "name": name, "debit": float(debit), "credit": float(credit)}
        for account, name, debit, credit in lines
    ]


# -------------------------------------------------------------------
# SEPA (pain.001.001.03)
# -------------------------------------------------------------------

def normalize_iban(value: str | None) -> str:
    return re.sub(r"\s+", "", value or "").upper()


def iban_is_valid(value: str | None) -> bool:
    iban = normalize_iban(value)

    if not re.fullmatch(r"[A-Z]{2}\d{2}[A-Z0-9]{10,30}", iban):
        return False

    rearranged = iban[4:] + iban[:4]
    digits = "".join(str(int(char, 36)) for char in rearranged)

    return int(digits) % 97 == 1


def sepa_text(value: str, limit: int = 70) -> str:
    # Juego de caracteres SEPA básico: sin tildes ni símbolos raros.
    import unicodedata

    normalized = unicodedata.normalize("NFKD", value)
    ascii_text = "".join(char for char in normalized if not unicodedata.combining(char))
    ascii_text = re.sub(r"[^A-Za-z0-9/\-?:().,'+ ]", " ", ascii_text)

    return escape(re.sub(r"\s+", " ", ascii_text).strip()[:limit])


def build_sepa_xml(
    run: PayrollRun,
    company: CompanyProfile | None,
    *,
    execution_date: date,
) -> tuple[bytes, list[str]]:
    if company is None or not company.iban or not iban_is_valid(company.iban):
        raise ValueError(
            "Indica un IBAN válido de la empresa en Mi empresa para generar la remesa."
        )

    payments = []
    skipped: list[str] = []

    for payslip in run.payslips:
        employee = payslip.employee

        if payslip.net <= 0:
            continue

        if employee is None or not iban_is_valid(employee.iban):
            skipped.append(payslip.employee_name)
            continue

        payments.append((payslip, normalize_iban(employee.iban)))

    if not payments:
        raise ValueError("Ninguna persona tiene un IBAN válido para pagar.")

    total = sum((money(payslip.net) for payslip, _iban in payments), ZERO)
    message_id = f"CF-NOM-{run.year}{run.month:02d}-{run.id}"
    created = clock.now().strftime("%Y-%m-%dT%H:%M:%S")
    debtor = sepa_text(company.name or "Empresa")
    bic = (company.bic or "").strip().upper()
    debtor_agent = (
        f"<DbtrAgt><FinInstnId><BIC>{escape(bic)}</BIC></FinInstnId></DbtrAgt>"
        if bic
        else "<DbtrAgt><FinInstnId><Othr><Id>NOTPROVIDED</Id></Othr></FinInstnId></DbtrAgt>"
    )
    concept = sepa_text(f"Nomina {period_label(run)}", 140)

    transactions = "".join(
        f"""
      <CdtTrfTxInf>
        <PmtId><EndToEndId>{escape(f"NOM{run.year}{run.month:02d}-{payslip.id}")}</EndToEndId></PmtId>
        <Amt><InstdAmt Ccy="EUR">{money(payslip.net)}</InstdAmt></Amt>
        <Cdtr><Nm>{sepa_text(payslip.employee_name)}</Nm></Cdtr>
        <CdtrAcct><Id><IBAN>{iban}</IBAN></Id></CdtrAcct>
        <Purp><Cd>SALA</Cd></Purp>
        <RmtInf><Ustrd>{concept}</Ustrd></RmtInf>
      </CdtTrfTxInf>"""
        for payslip, iban in payments
    )

    xml = f"""<?xml version="1.0" encoding="UTF-8"?>
<Document xmlns="urn:iso:std:iso:20022:tech:xsd:pain.001.001.03">
  <CstmrCdtTrfInitn>
    <GrpHdr>
      <MsgId>{message_id}</MsgId>
      <CreDtTm>{created}</CreDtTm>
      <NbOfTxs>{len(payments)}</NbOfTxs>
      <CtrlSum>{total}</CtrlSum>
      <InitgPty><Nm>{debtor}</Nm></InitgPty>
    </GrpHdr>
    <PmtInf>
      <PmtInfId>{message_id}-1</PmtInfId>
      <PmtMtd>TRF</PmtMtd>
      <NbOfTxs>{len(payments)}</NbOfTxs>
      <CtrlSum>{total}</CtrlSum>
      <PmtTpInf><SvcLvl><Cd>SEPA</Cd></SvcLvl><CtgyPurp><Cd>SALA</Cd></CtgyPurp></PmtTpInf>
      <ReqdExctnDt>{execution_date.isoformat()}</ReqdExctnDt>
      <Dbtr><Nm>{debtor}</Nm></Dbtr>
      <DbtrAcct><Id><IBAN>{normalize_iban(company.iban)}</IBAN></Id></DbtrAcct>
      {debtor_agent}
      <ChrgBr>SLEV</ChrgBr>{transactions}
    </PmtInf>
  </CstmrCdtTrfInitn>
</Document>
"""

    return xml.encode("utf-8"), skipped


# -------------------------------------------------------------------
# Recibos en PDF
# -------------------------------------------------------------------

def format_eur(value: Any) -> str:
    text = f"{float(value):,.2f}"
    return text.replace(",", "X").replace(".", ",").replace("X", ".") + " €"


def build_payslips_pdf(
    run: PayrollRun,
    company: CompanyProfile | None,
    payslips: list[Payslip] | None = None,
) -> bytes:
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas

    accent = HexColor("#8c1d33")
    muted = HexColor("#6b7280")
    line_color = HexColor("#e5e7eb")

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    items = payslips if payslips is not None else run.payslips

    for payslip in items:
        employee = payslip.employee
        lines = payslip.lines or {}
        y = height - 22 * mm

        pdf.setFillColor(accent)
        pdf.setFont("Helvetica-Bold", 15)
        pdf.drawString(20 * mm, y, "Recibo de salarios")
        pdf.setFont("Helvetica", 10)
        pdf.setFillColor(muted)
        pdf.drawRightString(width - 20 * mm, y, f"Periodo: {period_label(run)}")
        y -= 10 * mm

        pdf.setFillColor(HexColor("#111827"))
        pdf.setFont("Helvetica-Bold", 10)
        pdf.drawString(20 * mm, y, "Empresa")
        pdf.drawString(110 * mm, y, "Trabajador/a")
        pdf.setFont("Helvetica", 9.5)
        y -= 5 * mm
        company_lines = [
            company.name if company and company.name else "—",
            f"CIF: {company.tax_id}" if company and company.tax_id else "",
        ]
        employee_lines = [
            payslip.employee_name,
            f"NIF: {employee.tax_id}" if employee and employee.tax_id else "",
            f"Nº afiliación SS: {employee.ss_number}" if employee and employee.ss_number else "",
            f"Puesto: {employee.job_title}" if employee and employee.job_title else "",
        ]

        for index in range(max(len(company_lines), len(employee_lines))):
            if index < len(company_lines) and company_lines[index]:
                pdf.drawString(20 * mm, y, company_lines[index])
            if index < len(employee_lines) and employee_lines[index]:
                pdf.drawString(110 * mm, y, employee_lines[index])
            y -= 4.6 * mm

        y -= 4 * mm

        def section(title: str, rows: list[dict[str, Any]], total_label: str, total: Any) -> None:
            nonlocal y
            pdf.setFillColor(accent)
            pdf.setFont("Helvetica-Bold", 10)
            pdf.drawString(20 * mm, y, title)
            y -= 2 * mm
            pdf.setStrokeColor(line_color)
            pdf.line(20 * mm, y, width - 20 * mm, y)
            y -= 5 * mm
            pdf.setFillColor(HexColor("#111827"))
            pdf.setFont("Helvetica", 9.5)

            for row in rows:
                pdf.drawString(22 * mm, y, str(row["concept"]))
                pdf.drawRightString(width - 22 * mm, y, format_eur(row["amount"]))
                y -= 5 * mm

            pdf.setFont("Helvetica-Bold", 9.5)
            pdf.drawString(22 * mm, y, total_label)
            pdf.drawRightString(width - 22 * mm, y, format_eur(total))
            y -= 9 * mm

        section("I. Devengos", lines.get("earnings", []), "Total devengado", payslip.gross)
        section(
            "II. Deducciones",
            lines.get("deductions", []),
            "Total a deducir",
            money(payslip.ss_employee) + money(payslip.irpf) + money(payslip.advance),
        )

        pdf.setFillColor(HexColor("#f7ebed"))
        pdf.rect(20 * mm, y - 4 * mm, width - 40 * mm, 12 * mm, stroke=0, fill=1)
        pdf.setFillColor(accent)
        pdf.setFont("Helvetica-Bold", 12)
        pdf.drawString(24 * mm, y, "Líquido total a percibir (I − II)")
        pdf.drawRightString(width - 24 * mm, y, format_eur(payslip.net))
        y -= 16 * mm

        section(
            "Aportación de la empresa a la Seguridad Social",
            lines.get("employer", []),
            "Total aportación empresarial",
            payslip.ss_employer,
        )

        base = lines.get("base", {})
        pdf.setFont("Helvetica", 8.5)
        pdf.setFillColor(muted)
        pdf.drawString(
            20 * mm,
            y,
            f"Base de cotización aplicada: {format_eur(base.get('applied', payslip.contribution_base))}"
            f" · Coste total empresa: {format_eur(payslip.company_cost)}",
        )
        y -= 4.5 * mm
        pdf.drawString(20 * mm, y, f"Parámetros: {lines.get('params_source', '')}")
        y -= 4.5 * mm
        pdf.drawString(
            20 * mm,
            y,
            "Documento generado por CapaFiscal como borrador. Debe revisarse antes de su entrega.",
        )

        pdf.setFont("Helvetica", 9)
        pdf.setFillColor(HexColor("#111827"))
        pdf.drawString(20 * mm, 30 * mm, "Firma y sello de la empresa")
        pdf.drawString(120 * mm, 30 * mm, "Recibí")
        pdf.showPage()

    pdf.save()

    return buffer.getvalue()


# -------------------------------------------------------------------
# Resumen Excel para la gestoría
# -------------------------------------------------------------------

def build_summary_xlsx(run: PayrollRun) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.styles import PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Nóminas"
    sheet.append([f"Resumen de nóminas · {period_label(run)}"])
    sheet["A1"].font = Font(bold=True, size=13)
    sheet.append([])
    headers = [
        "Persona", "Días", "Devengado", "Base cotización", "SS trabajador",
        "IRPF %", "IRPF", "Anticipos", "Líquido", "SS empresa", "Coste empresa",
    ]
    sheet.append(headers)

    for cell in sheet[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="8C1D33")

    for payslip in run.payslips:
        sheet.append(
            [
                payslip.employee_name,
                (payslip.lines or {}).get("days", 30),
                float(payslip.gross),
                float(payslip.contribution_base),
                float(payslip.ss_employee),
                float(payslip.irpf_rate),
                float(payslip.irpf),
                float(payslip.advance),
                float(payslip.net),
                float(payslip.ss_employer),
                float(payslip.company_cost),
            ]
        )

    totals = run_totals(run)
    sheet.append(
        ["TOTAL", "", totals["gross"], "", totals["ss_employee"], "", totals["irpf"], "",
         totals["net"], totals["ss_employer"], totals["company_cost"]]
    )

    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)

    for column in "CDEGHIJK":
        for cell in sheet[column][3:]:
            cell.number_format = '#,##0.00 "€"'

    sheet.column_dimensions["A"].width = 32

    for column in "BCDEFGHIJK":
        sheet.column_dimensions[column].width = 15

    journal = workbook.create_sheet("Asiento")
    journal.append(["Cuenta", "Descripción", "Debe", "Haber"])

    for cell in journal[1]:
        cell.font = Font(bold=True)

    for line in journal_entry(run):
        journal.append([line["account"], line["name"], line["debit"] or None, line["credit"] or None])

    journal.column_dimensions["B"].width = 48

    output = io.BytesIO()
    workbook.save(output)

    return output.getvalue()
