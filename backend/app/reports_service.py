from __future__ import annotations

import csv
import io
from collections import defaultdict
from datetime import date
from decimal import Decimal
from decimal import ROUND_HALF_UP
from typing import Any

from sqlalchemy import or_
from sqlalchemy import select
from sqlalchemy.orm import Session
from sqlalchemy.orm import selectinload

from app.extractor import category_account
from app.extractor import is_valid_spanish_tax_id
from app.models import Document
from app.models import Invoice


ZERO = Decimal("0.00")
CENT = Decimal("0.01")

DUE_SOON_DAYS = 7


# -------------------------------------------------------------------
# Utilidades
# -------------------------------------------------------------------

RECEIVED = "RECEIVED"
ISSUED = "ISSUED"


def normalize_direction(value: str | None) -> str:
    return ISSUED if (value or "").strip().upper() in {"ISSUED", "EMITIDAS"} else RECEIVED


def direction_clause(direction: str):
    # Las facturas antiguas (sin sentido) cuentan como recibidas.
    if direction == ISSUED:
        return Invoice.direction == ISSUED

    return or_(Invoice.direction.is_(None), Invoice.direction != ISSUED)


def is_issued(invoice: Invoice) -> bool:
    return invoice.direction == ISSUED


def counterparty(invoice: Invoice) -> tuple[str | None, str | None]:
    if is_issued(invoice):
        return invoice.customer_name, invoice.customer_tax_id

    return invoice.supplier_name, invoice.supplier_tax_id


def money(value: Decimal | None) -> Decimal:
    if value is None:
        return ZERO

    return Decimal(value).quantize(CENT, rounding=ROUND_HALF_UP)


def to_number(value: Decimal | None) -> float:
    return float(money(value))


def quarter_range(
    year: int,
    quarter: int | None,
) -> tuple[date, date]:
    if quarter is None:
        return date(year, 1, 1), date(year, 12, 31)

    if quarter not in {1, 2, 3, 4}:
        raise ValueError("El trimestre debe estar entre 1 y 4.")

    start_month = (quarter - 1) * 3 + 1
    end_month = start_month + 2
    end_day = 31 if end_month in {3, 12} else 30

    return date(year, start_month, 1), date(year, end_month, end_day)


def period_label(year: int, quarter: int | None) -> str:
    if quarter is None:
        return f"Ejercicio {year}"

    return f"{quarter}T {year}"


def load_invoices_in_period(
    database: Session,
    *,
    date_from: date,
    date_to: date,
    review_statuses: set[str] | None = None,
    direction: str | None = None,
) -> list[Invoice]:
    statement = (
        select(Invoice)
        .where(
            Invoice.invoice_date >= date_from,
            Invoice.invoice_date <= date_to,
        )
        .options(
            selectinload(Invoice.tax_lines),
            selectinload(Invoice.document),
        )
        .order_by(
            Invoice.invoice_date.asc(),
            Invoice.id.asc(),
        )
    )

    if review_statuses:
        statement = statement.where(
            Invoice.review_status.in_(review_statuses)
        )

    if direction:
        statement = statement.where(direction_clause(direction))

    return list(database.scalars(statement).all())


def invoice_tax_breakdown(
    invoice: Invoice,
) -> list[dict[str, Decimal | None]]:
    """
    Desglose por tipo impositivo. Si la factura no tiene líneas de
    impuesto, se deduce una línea única a partir de base y cuota.
    """
    lines = [
        {
            "tax_rate": line.tax_rate,
            "tax_base": line.tax_base,
            "tax_amount": line.tax_amount,
        }
        for line in invoice.tax_lines
        if line.tax_base is not None or line.tax_amount is not None
    ]

    if lines:
        return lines

    rate: Decimal | None = None

    if invoice.subtotal and invoice.tax_total is not None:
        rate = (
            Decimal(invoice.tax_total)
            / Decimal(invoice.subtotal)
            * 100
        ).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)

        # Redondea a los tipos españoles habituales si está muy cerca.
        for standard_rate in (Decimal("21"), Decimal("10"), Decimal("5"),
                              Decimal("4"), Decimal("0")):
            if abs(rate - standard_rate) <= Decimal("0.3"):
                rate = standard_rate
                break

    return [
        {
            "tax_rate": rate,
            "tax_base": invoice.subtotal,
            "tax_amount": invoice.tax_total,
        }
    ]


def normalize_rate(rate: Decimal | float | str) -> Decimal:
    # 21.000 -> 21 ; 5.200 -> 5.2 (sin notación exponencial).
    normalized = Decimal(str(rate)).normalize()

    if normalized == normalized.to_integral():
        return normalized.quantize(Decimal(1))

    return normalized


def format_eur(value: Decimal | float) -> str:
    text = f"{float(value):,.2f}"

    return text.replace(",", "X").replace(".", ",").replace("X", ".") + " €"


def rate_label(rate: Decimal | None) -> str:
    if rate is None:
        return "Sin tipo"

    text = format(normalize_rate(rate), "f").replace(".", ",")

    return f"{text} %"


# -------------------------------------------------------------------
# Resumen fiscal del periodo
# -------------------------------------------------------------------

def summarize_by_rate(invoices: list[Invoice]) -> dict[str, Any]:
    by_rate: dict[str, dict[str, Any]] = {}
    base = tax = withholding = total = ZERO

    for invoice in invoices:
        base += money(invoice.subtotal)
        tax += money(invoice.tax_total)
        withholding += money(invoice.withholding_total)
        total += money(invoice.total)

        for line in invoice_tax_breakdown(invoice):
            key = rate_label(line["tax_rate"])
            bucket = by_rate.setdefault(
                key,
                {
                    "rate": (
                        float(line["tax_rate"])
                        if line["tax_rate"] is not None
                        else None
                    ),
                    "label": key,
                    "base": ZERO,
                    "tax": ZERO,
                    "invoices": 0,
                },
            )
            bucket["base"] += money(line["tax_base"])
            bucket["tax"] += money(line["tax_amount"])
            bucket["invoices"] += 1

    return {
        "base": float(base),
        "tax": float(tax),
        "withholding": float(withholding),
        "total": float(total),
        "by_rate": sorted(
            (
                {
                    **bucket,
                    "base": float(bucket["base"]),
                    "tax": float(bucket["tax"]),
                }
                for bucket in by_rate.values()
            ),
            key=lambda item: -(item["rate"] if item["rate"] is not None else -1),
        ),
    }


def build_vat_report(
    database: Session,
    *,
    year: int,
    quarter: int | None,
) -> dict[str, Any]:
    date_from, date_to = quarter_range(year, quarter)

    invoices = load_invoices_in_period(
        database,
        date_from=date_from,
        date_to=date_to,
    )

    approved = [
        invoice
        for invoice in invoices
        if invoice.review_status == "APPROVED"
        and not is_issued(invoice)
    ]
    approved_issued = [
        invoice
        for invoice in invoices
        if invoice.review_status == "APPROVED"
        and is_issued(invoice)
    ]
    pending = [
        invoice
        for invoice in invoices
        if invoice.review_status == "PENDING"
    ]

    issued_summary = summarize_by_rate(approved_issued)

    by_rate: dict[str, dict[str, Any]] = {}
    by_category: dict[str, dict[str, Any]] = {}
    by_month: dict[str, dict[str, Any]] = {}

    total_base = ZERO
    total_tax = ZERO
    total_withholding = ZERO
    total_surcharge = ZERO
    total_amount = ZERO

    for invoice in approved:
        total_base += money(invoice.subtotal)
        total_tax += money(invoice.tax_total)
        total_withholding += money(invoice.withholding_total)
        total_surcharge += money(invoice.surcharge_total)
        total_amount += money(invoice.total)

        for line in invoice_tax_breakdown(invoice):
            key = rate_label(line["tax_rate"])
            bucket = by_rate.setdefault(
                key,
                {
                    "rate": (
                        float(line["tax_rate"])
                        if line["tax_rate"] is not None
                        else None
                    ),
                    "label": key,
                    "base": ZERO,
                    "tax": ZERO,
                    "invoices": 0,
                },
            )
            bucket["base"] += money(line["tax_base"])
            bucket["tax"] += money(line["tax_amount"])
            bucket["invoices"] += 1

        category = invoice.category or "Otros gastos"
        category_bucket = by_category.setdefault(
            category,
            {
                "category": category,
                "account": category_account(category),
                "base": ZERO,
                "tax": ZERO,
                "total": ZERO,
                "invoices": 0,
            },
        )
        category_bucket["base"] += money(invoice.subtotal)
        category_bucket["tax"] += money(invoice.tax_total)
        category_bucket["total"] += money(invoice.total)
        category_bucket["invoices"] += 1

        month_key = invoice.invoice_date.strftime("%Y-%m")
        month_bucket = by_month.setdefault(
            month_key,
            {
                "month": month_key,
                "base": ZERO,
                "tax": ZERO,
                "total": ZERO,
                "invoices": 0,
            },
        )
        month_bucket["base"] += money(invoice.subtotal)
        month_bucket["tax"] += money(invoice.tax_total)
        month_bucket["total"] += money(invoice.total)
        month_bucket["invoices"] += 1

    def serialize(bucket: dict[str, Any]) -> dict[str, Any]:
        return {
            key: (float(value) if isinstance(value, Decimal) else value)
            for key, value in bucket.items()
        }

    warnings: list[str] = []

    if pending:
        pending_amount = sum(
            (money(invoice.total) for invoice in pending),
            ZERO,
        )
        warnings.append(
            f"Hay {len(pending)} factura(s) del periodo pendientes de "
            f"revisión por {format_eur(pending_amount)} que no se "
            "incluyen en este resumen."
        )

    undated_statement = (
        select(Invoice.id)
        .where(
            Invoice.invoice_date.is_(None),
            Invoice.review_status != "REJECTED",
        )
    )
    undated_count = len(database.scalars(undated_statement).all())

    if undated_count:
        warnings.append(
            f"{undated_count} factura(s) no tienen fecha y no pueden "
            "asignarse a ningún periodo."
        )

    return {
        "period": period_label(year, quarter),
        "year": year,
        "quarter": quarter,
        "date_from": date_from.isoformat(),
        "date_to": date_to.isoformat(),
        "approved_invoices": len(approved),
        "pending_invoices": len(pending),
        "total_base": float(total_base),
        "total_tax": float(total_tax),
        "total_withholding": float(total_withholding),
        "total_surcharge": float(total_surcharge),
        "total_amount": float(total_amount),
        "by_rate": sorted(
            (serialize(bucket) for bucket in by_rate.values()),
            key=lambda item: -(item["rate"] or -1),
        ),
        "by_category": sorted(
            (serialize(bucket) for bucket in by_category.values()),
            key=lambda item: -item["total"],
        ),
        "by_month": sorted(
            (serialize(bucket) for bucket in by_month.values()),
            key=lambda item: item["month"],
        ),
        "issued_invoices": len(approved_issued),
        "issued_base": issued_summary["base"],
        "issued_tax": issued_summary["tax"],
        "issued_withholding": issued_summary["withholding"],
        "issued_total": issued_summary["total"],
        "issued_by_rate": issued_summary["by_rate"],
        "vat_balance": round(issued_summary["tax"] - float(total_tax), 2),
        "warnings": warnings,
        "note": (
            "IVA soportado de facturas recibidas aprobadas, agrupado por "
            "fecha de factura. Es una ayuda para preparar el modelo 303 "
            "(y el 111 si hay retenciones), no una autoliquidación."
        ),
    }


# -------------------------------------------------------------------
# Proveedores
# -------------------------------------------------------------------

def supplier_key(invoice: Invoice) -> str | None:
    name, tax_id = counterparty(invoice)

    if tax_id:
        return f"nif:{tax_id}"

    if name:
        return f"name:{name.strip().upper()}"

    return None


def build_supplier_list(
    database: Session,
    *,
    search: str | None = None,
    direction: str = RECEIVED,
) -> list[dict[str, Any]]:
    statement = (
        select(Invoice)
        .where(direction_clause(direction))
        .options(selectinload(Invoice.document))
    )
    invoices = list(database.scalars(statement).all())

    suppliers: dict[str, dict[str, Any]] = {}
    categories: dict[str, dict[str, int]] = defaultdict(
        lambda: defaultdict(int)
    )

    for invoice in invoices:
        key = supplier_key(invoice)

        if key is None or invoice.review_status == "REJECTED":
            continue

        party_name, party_tax_id = counterparty(invoice)

        supplier = suppliers.setdefault(
            key,
            {
                "key": key,
                "name": party_name,
                "tax_id": party_tax_id,
                "tax_id_valid": is_valid_spanish_tax_id(party_tax_id),
                "invoices": 0,
                "approved_invoices": 0,
                "pending_invoices": 0,
                "total_approved": ZERO,
                "tax_approved": ZERO,
                "unpaid_amount": ZERO,
                "last_invoice_date": None,
                "main_category": None,
            },
        )

        if not supplier["name"] and party_name:
            supplier["name"] = party_name

        supplier["invoices"] += 1

        if invoice.review_status == "APPROVED":
            supplier["approved_invoices"] += 1
            supplier["total_approved"] += money(invoice.total)
            supplier["tax_approved"] += money(invoice.tax_total)

            if invoice.paid_at is None:
                supplier["unpaid_amount"] += money(invoice.total)
        else:
            supplier["pending_invoices"] += 1

        if invoice.invoice_date and (
            supplier["last_invoice_date"] is None
            or invoice.invoice_date > supplier["last_invoice_date"]
        ):
            supplier["last_invoice_date"] = invoice.invoice_date

        if invoice.category:
            categories[key][invoice.category] += 1

    normalized_search = (search or "").strip().lower()
    result: list[dict[str, Any]] = []

    for key, supplier in suppliers.items():
        if categories[key]:
            supplier["main_category"] = max(
                categories[key].items(),
                key=lambda item: item[1],
            )[0]

        if normalized_search and not (
            normalized_search in (supplier["name"] or "").lower()
            or normalized_search in (supplier["tax_id"] or "").lower()
        ):
            continue

        result.append(
            {
                **supplier,
                "total_approved": float(supplier["total_approved"]),
                "tax_approved": float(supplier["tax_approved"]),
                "unpaid_amount": float(supplier["unpaid_amount"]),
                "last_invoice_date": (
                    supplier["last_invoice_date"].isoformat()
                    if supplier["last_invoice_date"]
                    else None
                ),
            }
        )

    result.sort(key=lambda item: -item["total_approved"])

    return result


# -------------------------------------------------------------------
# Pagos y vencimientos
# -------------------------------------------------------------------

def payment_state(
    invoice: Invoice,
    today: date,
) -> str:
    if invoice.paid_at is not None:
        return "PAID"

    if invoice.due_date is None:
        return "NO_DUE_DATE"

    if invoice.due_date < today:
        return "OVERDUE"

    if (invoice.due_date - today).days <= DUE_SOON_DAYS:
        return "DUE_SOON"

    return "PENDING"


def build_payments_overview(
    database: Session,
    *,
    today: date | None = None,
    direction: str = RECEIVED,
) -> dict[str, Any]:
    """
    Recibidas: pagos pendientes a proveedores.
    Emitidas: cobros pendientes de clientes (paid_at = fecha de cobro).
    """
    current_day = today or date.today()

    statement = (
        select(Invoice)
        .where(
            Invoice.review_status == "APPROVED",
            Invoice.paid_at.is_(None),
            direction_clause(direction),
        )
        .options(selectinload(Invoice.document))
    )
    invoices = list(database.scalars(statement).all())

    items: list[dict[str, Any]] = []
    totals: dict[str, Decimal] = defaultdict(lambda: ZERO)
    counts: dict[str, int] = defaultdict(int)

    for invoice in invoices:
        state = payment_state(invoice, current_day)
        totals[state] += money(invoice.total)
        counts[state] += 1
        party_name, party_tax_id = counterparty(invoice)

        items.append(
            {
                "invoice_id": invoice.id,
                "document_id": invoice.document_id,
                "direction": direction,
                "supplier_name": party_name,
                "supplier_tax_id": party_tax_id,
                "invoice_number": invoice.invoice_number,
                "invoice_date": (
                    invoice.invoice_date.isoformat()
                    if invoice.invoice_date
                    else None
                ),
                "due_date": (
                    invoice.due_date.isoformat()
                    if invoice.due_date
                    else None
                ),
                "days_to_due": (
                    (invoice.due_date - current_day).days
                    if invoice.due_date
                    else None
                ),
                "total": to_number(invoice.total),
                "currency": invoice.currency,
                "state": state,
            }
        )

    state_order = {
        "OVERDUE": 0,
        "DUE_SOON": 1,
        "PENDING": 2,
        "NO_DUE_DATE": 3,
    }

    items.sort(
        key=lambda item: (
            state_order.get(item["state"], 9),
            item["due_date"] or "9999-12-31",
        )
    )

    return {
        "generated_for": current_day.isoformat(),
        "direction": direction,
        "unpaid_count": len(items),
        "unpaid_total": float(sum(totals.values(), ZERO)),
        "overdue_count": counts["OVERDUE"],
        "overdue_total": float(totals["OVERDUE"]),
        "due_soon_count": counts["DUE_SOON"],
        "due_soon_total": float(totals["DUE_SOON"]),
        "items": items,
    }


# -------------------------------------------------------------------
# Libro registro de facturas recibidas
# -------------------------------------------------------------------

ISSUED_LEDGER_LABELS = {
    "supplier_tax_id": "NIF destinatario",
    "supplier_name": "Nombre destinatario",
    "paid_at": "Fecha cobro",
    "withholding": "Retención IRPF soportada",
    "surcharge": "Recargo equivalencia",
}


def ledger_columns(direction: str = RECEIVED) -> tuple[tuple[str, str], ...]:
    if direction != ISSUED:
        return LEDGER_COLUMNS

    return tuple(
        (key, ISSUED_LEDGER_LABELS.get(key, label))
        for key, label in LEDGER_COLUMNS
    )


LEDGER_COLUMNS = (
    ("order", "Nº orden"),
    ("invoice_date", "Fecha expedición"),
    ("registered_at", "Fecha registro"),
    ("invoice_number", "Nº factura"),
    ("supplier_tax_id", "NIF emisor"),
    ("supplier_name", "Nombre emisor"),
    ("concept", "Concepto"),
    ("category", "Categoría"),
    ("account", "Cuenta PGC"),
    ("tax_base", "Base imponible"),
    ("tax_rate", "Tipo IVA %"),
    ("tax_amount", "Cuota IVA"),
    ("surcharge", "Recargo equivalencia"),
    ("withholding", "Retención IRPF"),
    ("total", "Total factura"),
    ("paid_at", "Fecha pago"),
    ("payment_method", "Forma de pago"),
    ("document_id", "ID documento"),
)


def build_ledger_rows(
    database: Session,
    *,
    date_from: date,
    date_to: date,
    include_pending: bool = False,
    direction: str = RECEIVED,
) -> list[dict[str, Any]]:
    statuses = {"APPROVED"}

    if include_pending:
        statuses.add("PENDING")

    invoices = load_invoices_in_period(
        database,
        date_from=date_from,
        date_to=date_to,
        review_statuses=statuses,
        direction=direction,
    )

    rows: list[dict[str, Any]] = []

    for order, invoice in enumerate(invoices, start=1):
        breakdown = invoice_tax_breakdown(invoice)

        for line_index, line in enumerate(breakdown):
            # Retención, recargo y total solo en la primera línea para
            # que las sumas por columna no se dupliquen.
            first_line = line_index == 0

            rows.append(
                {
                    "order": order,
                    "invoice_date": invoice.invoice_date,
                    "registered_at": (
                        invoice.approved_at.date()
                        if invoice.approved_at
                        else None
                    ),
                    "invoice_number": invoice.invoice_number,
                    "supplier_tax_id": counterparty(invoice)[1],
                    "supplier_name": counterparty(invoice)[0],
                    "concept": (invoice.concept or "")[:250],
                    "category": invoice.category,
                    "account": category_account(invoice.category),
                    "tax_base": money(line["tax_base"]),
                    "tax_rate": (
                        normalize_rate(line["tax_rate"])
                        if line["tax_rate"] is not None
                        else None
                    ),
                    "tax_amount": money(line["tax_amount"]),
                    "surcharge": (
                        money(invoice.surcharge_total)
                        if first_line
                        else ZERO
                    ),
                    "withholding": (
                        money(invoice.withholding_total)
                        if first_line
                        else ZERO
                    ),
                    "total": money(invoice.total) if first_line else ZERO,
                    "paid_at": invoice.paid_at,
                    "payment_method": invoice.payment_method,
                    "document_id": invoice.document_id,
                }
            )

    return rows


def format_csv_value(value: Any) -> str:
    if value is None:
        return ""

    if isinstance(value, Decimal):
        return format(value, "f").replace(".", ",")

    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")

    return str(value)


def ledger_to_csv(
    rows: list[dict[str, Any]],
    direction: str = RECEIVED,
) -> bytes:
    columns = ledger_columns(direction)
    buffer = io.StringIO()
    writer = csv.writer(
        buffer,
        delimiter=";",
        quoting=csv.QUOTE_MINIMAL,
    )

    writer.writerow(label for _key, label in columns)

    for row in rows:
        writer.writerow(
            format_csv_value(row[key])
            for key, _label in columns
        )

    # BOM para que Excel en español detecte UTF-8.
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def ledger_to_xlsx(
    rows: list[dict[str, Any]],
    *,
    title: str,
    direction: str = RECEIVED,
) -> bytes:
    columns = ledger_columns(direction)

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.styles import PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Facturas expedidas" if direction == ISSUED else "Facturas recibidas"

    sheet.append([title])
    sheet["A1"].font = Font(bold=True, size=13)
    sheet.append([])

    header_row = 3
    sheet.append([label for _key, label in columns])

    header_fill = PatternFill("solid", fgColor="9E1B32")

    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    money_keys = {
        "tax_base", "tax_amount", "surcharge", "withholding", "total",
    }
    date_keys = {"invoice_date", "registered_at", "paid_at"}

    for row in rows:
        values = []

        for key, _label in columns:
            value = row[key]

            if isinstance(value, Decimal):
                value = float(value)

            values.append(value)

        sheet.append(values)

    for column_index, (key, label) in enumerate(columns, start=1):
        letter = get_column_letter(column_index)

        if key in money_keys:
            for cell in sheet[letter][header_row:]:
                cell.number_format = '#,##0.00 "€"'
        elif key in date_keys:
            for cell in sheet[letter][header_row:]:
                cell.number_format = "DD/MM/YYYY"

        width = 14

        if key in {"supplier_name", "concept"}:
            width = 38
        elif key == "category":
            width = 26

        sheet.column_dimensions[letter].width = max(width, len(label) + 2)

    if rows:
        total_row = sheet.max_row + 1
        sheet.cell(row=total_row, column=1, value="TOTALES").font = Font(
            bold=True
        )

        for column_index, (key, _label) in enumerate(
            columns,
            start=1,
        ):
            if key not in money_keys:
                continue

            letter = get_column_letter(column_index)
            cell = sheet.cell(
                row=total_row,
                column=column_index,
                value=(
                    f"=SUM({letter}{header_row + 1}:"
                    f"{letter}{total_row - 1})"
                ),
            )
            cell.font = Font(bold=True)
            cell.number_format = '#,##0.00 "€"'

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    output = io.BytesIO()
    workbook.save(output)

    return output.getvalue()


# -------------------------------------------------------------------
# Búsqueda de documentos
# -------------------------------------------------------------------

def apply_document_filters(
    statement,
    *,
    search: str | None,
    review_status: str | None,
    category: str | None,
    date_from: date | None,
    date_to: date | None,
    payment: str | None,
    direction: str | None = None,
    today: date | None = None,
):
    current_day = today or date.today()
    needs_invoice = any(
        (search, review_status, category, date_from, date_to, payment,
         direction)
    )

    if not needs_invoice:
        return statement

    statement = statement.outerjoin(
        Invoice,
        Invoice.document_id == Document.id,
    )

    if direction:
        statement = statement.where(
            Invoice.id.is_not(None),
            direction_clause(normalize_direction(direction)),
        )

    if search:
        pattern = f"%{search.strip()}%"
        statement = statement.where(
            Document.original_filename.ilike(pattern)
            | Invoice.supplier_name.ilike(pattern)
            | Invoice.supplier_tax_id.ilike(pattern)
            | Invoice.customer_name.ilike(pattern)
            | Invoice.customer_tax_id.ilike(pattern)
            | Invoice.invoice_number.ilike(pattern)
            | Invoice.concept.ilike(pattern)
        )

    if review_status:
        statement = statement.where(
            Invoice.review_status == review_status.strip().upper()
        )

    if category:
        statement = statement.where(Invoice.category == category.strip())

    if date_from:
        statement = statement.where(Invoice.invoice_date >= date_from)

    if date_to:
        statement = statement.where(Invoice.invoice_date <= date_to)

    if payment:
        normalized_payment = payment.strip().lower()

        if normalized_payment == "paid":
            statement = statement.where(Invoice.paid_at.is_not(None))
        elif normalized_payment == "unpaid":
            statement = statement.where(
                Invoice.review_status == "APPROVED",
                Invoice.paid_at.is_(None),
            )
        elif normalized_payment == "overdue":
            statement = statement.where(
                Invoice.review_status == "APPROVED",
                Invoice.paid_at.is_(None),
                Invoice.due_date < current_day,
            )

    return statement
