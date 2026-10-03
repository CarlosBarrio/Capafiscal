"""
Contabilización (BORRADOR): el libro diario sale de lo que ya está en CapaFiscal, sin volver a teclear nada.

    factura recibida aprobada   6xx (base, + recargo) + 472 IVA soportado  a  400/410 proveedor (+ 4751 retención)
    factura emitida aprobada    430 cliente (+ 473 retención)              a  7xx (base) + 477 IVA repercutido
    pago / cobro conciliado     400/410  a  572   ·   572  a  430
    movimiento justificado      626 comisión · 465 nóminas · 476 Seguridad Social · 4750 Hacienda
                                438 anticipo · 555 traspasos y devoluciones (se anulan entre sí)

Solo lee. Cada asiento cuadra (debe = haber) o se marca. Lo que no se puede contabilizar
(factura sin aprobar, movimiento sin conciliar) se dice, no se inventa.

Exportación: libro diario genérico (CSV / Excel) que importa cualquier programa contable
con una plantilla de columnas. Los formatos propios de A3, Sage o Holded necesitan su
especificación de importación: no se imitan sin ella.
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session
from sqlalchemy.orm import selectinload

from app.models import BankAllocation
from app.models import BankTransaction
from app.models import Invoice

CENT = Decimal("0.01")
ACCOUNTS = {
    "400": "Proveedores", "410": "Acreedores por prestaciones de servicios", "430": "Clientes", "438": "Anticipos de clientes",
    "465": "Remuneraciones pendientes de pago", "472": "H.P. IVA soportado", "473": "H.P. retenciones y pagos a cuenta",
    "4750": "H.P. acreedora por conceptos fiscales", "4751": "H.P. acreedora por retenciones practicadas", "476": "Organismos de la S.S. acreedores",
    "477": "H.P. IVA repercutido", "555": "Partidas pendientes de aplicación", "572": "Bancos e instituciones de crédito c/c",
    "626": "Servicios bancarios y similares",
}
JUSTIFIED_ACCOUNT = {"COMISION": "626", "NOMINA": "465", "SEG_SOCIAL": "476", "IMPUESTO": "4750", "ANTICIPO": "438",
                     "TRASPASO": "555", "DEVOLUCION": "555", "OTRO": "555"}
COLUMNS = ("asiento", "fecha", "cuenta", "nombre_cuenta", "nif", "tercero", "concepto", "debe", "haber", "documento")


def money(value: Any) -> Decimal:
    return Decimal(str(value or 0)).quantize(CENT)


def line(account: str, debit: Decimal = Decimal("0"), credit: Decimal = Decimal("0"), *, name: str | None = None,
         tax_id: str | None = None, party: str | None = None) -> dict[str, Any]:
    debit, credit = money(debit), money(credit)
    if debit < 0 or credit < 0:  # rectificativas: un importe negativo va al lado contrario
        debit, credit = max(debit, Decimal("0")) - min(credit, Decimal("0")), max(credit, Decimal("0")) - min(debit, Decimal("0"))
    return {"account": account, "account_name": name or ACCOUNTS.get(account, ""), "tax_id": tax_id, "party": party,
            "debit": debit, "credit": credit}


def entry(when: date, concept: str, lines: list[dict[str, Any]], *, source: dict[str, Any], document: str | None) -> dict[str, Any]:
    lines = [item for item in lines if item["debit"] or item["credit"]]
    debit, credit = sum(item["debit"] for item in lines), sum(item["credit"] for item in lines)
    return {"date": when, "concept": concept[:120], "lines": lines, "source": source, "document": document,
            "debit": debit, "credit": credit, "balanced": abs(debit - credit) < CENT}


def party_account(invoice: Invoice, expense_account: str) -> str:
    """400 si se compran bienes (grupo 60), 410 si son servicios."""
    return "400" if expense_account.startswith("60") else "410"


def invoice_entry(invoice: Invoice) -> dict[str, Any]:
    from app.extractor import category_account
    from app.reports_service import counterparty
    from app.reports_service import invoice_tax_breakdown
    from app.reports_service import is_issued

    name, tax_id = counterparty(invoice)
    account = category_account(invoice.category)
    total = money(invoice.total)  # con signo: las rectificativas restan
    vat = sum((money(item["tax_amount"]) for item in invoice_tax_breakdown(invoice)), Decimal("0"))
    surcharge, withholding = money(invoice.surcharge_total), money(invoice.withholding_total)
    base = money(invoice.subtotal) if invoice.subtotal is not None else total - vat - surcharge + withholding
    number = invoice.invoice_number or "s/n"
    source = {"type": "invoice", "id": invoice.id, "document_id": invoice.document_id}
    if is_issued(invoice):
        income = account if account.startswith("7") else "705"
        return entry(invoice.invoice_date, f"Factura emitida {number} · {name or ''}", [
            line("430", debit=total, tax_id=tax_id, party=name), line("473", debit=withholding),
            line(income, credit=base, name=invoice.category or ACCOUNTS.get(income)), line("477", credit=vat + surcharge),
        ], source=source, document=number)
    creditor = party_account(invoice, account)
    return entry(invoice.invoice_date, f"Factura recibida {number} · {name or ''}", [
        line(account, debit=base + surcharge, name=invoice.category), line("472", debit=vat),
        line(creditor, credit=total, tax_id=tax_id, party=name), line("4751", credit=withholding),
    ], source=source, document=number)


def bank_entries(database: Session, date_from: date, date_to: date) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    from app.extractor import category_account
    from app.reports_service import counterparty
    from app.reports_service import is_issued

    transactions = database.scalars(select(BankTransaction).where(BankTransaction.booking_date.between(date_from, date_to))
                                    .order_by(BankTransaction.booking_date, BankTransaction.id)).all()
    allocations: dict[int, list[BankAllocation]] = {}
    if transactions:
        for item in database.scalars(select(BankAllocation).where(BankAllocation.transaction_id.in_([row.id for row in transactions]))).all():
            allocations.setdefault(item.transaction_id, []).append(item)
    entries, pending = [], []
    for transaction in transactions:
        inflow = transaction.amount > 0
        bank = line("572", name=f"Bancos · {transaction.account_label}" if transaction.account_label else None)
        parts: list[tuple[str, Decimal, Invoice | None, str]] = []
        if allocations.get(transaction.id):
            for item in allocations[transaction.id]:
                invoice = database.get(Invoice, item.invoice_id) if item.invoice_id else None
                parts.append((item.kind, money(item.amount), invoice, item.note or ""))
        elif transaction.match_status == "MATCHED" and transaction.matched_invoice_id:
            parts.append(("FACTURA", abs(money(transaction.amount)), database.get(Invoice, transaction.matched_invoice_id), ""))
        if not parts:
            if transaction.match_status != "IGNORED":
                pending.append({"date": transaction.booking_date.isoformat(), "description": transaction.description, "amount": float(transaction.amount),
                                "why": "sin conciliar: no se contabiliza hasta saber qué es"})
            continue
        lines = []
        for kind, amount, invoice, _note in parts:
            if kind == "FACTURA" and invoice is not None:
                name, tax_id = counterparty(invoice)
                account = "430" if is_issued(invoice) else party_account(invoice, category_account(invoice.category))
                lines.append(line(account, debit=Decimal("0") if inflow else amount, credit=amount if inflow else Decimal("0"), tax_id=tax_id, party=name))
            else:
                account = JUSTIFIED_ACCOUNT.get(kind, "555")
                lines.append(line(account, debit=Decimal("0") if inflow else amount, credit=amount if inflow else Decimal("0")))
        applied = sum(amount for _kind, amount, _invoice, _note in parts)
        bank.update(debit=applied if inflow else Decimal("0"), credit=Decimal("0") if inflow else applied)
        entries.append(entry(transaction.booking_date, f"{'Cobro' if inflow else 'Pago'} · {transaction.description}", [*lines, bank] if not inflow else [bank, *lines],
                             source={"type": "bank", "id": transaction.id}, document=None))
    return entries, pending


def journal(database: Session, date_from: date, date_to: date) -> dict[str, Any]:
    """Libro diario del periodo con sus comprobaciones. Solo lee."""
    invoices = database.scalars(select(Invoice).where(Invoice.invoice_date.between(date_from, date_to), Invoice.total.is_not(None))
                                .options(selectinload(Invoice.tax_lines)).order_by(Invoice.invoice_date, Invoice.id)).all()
    entries = [invoice_entry(invoice) for invoice in invoices if invoice.review_status == "APPROVED"]
    not_booked = [{"date": invoice.invoice_date.isoformat() if invoice.invoice_date else None, "description": f"Factura {invoice.invoice_number or 's/n'}",
                   "amount": float(invoice.total), "why": "sin aprobar: se contabiliza al aprobarla"}
                  for invoice in invoices if invoice.review_status == "PENDING"]
    movements, pending = bank_entries(database, date_from, date_to)
    entries = sorted(entries + movements, key=lambda item: (item["date"], 0 if item["source"]["type"] == "invoice" else 1))
    for number, item in enumerate(entries, start=1):
        item["number"] = number

    from app.extractor import DEFAULT_ACCOUNT

    unbalanced = [item["number"] for item in entries if not item["balanced"]]
    default_account = sum(1 for item in entries if any(row["account"] == DEFAULT_ACCOUNT for row in item["lines"]))
    no_tax_id = sum(1 for item in entries if any(row["account"] in ("400", "410", "430") and not row["tax_id"] for row in item["lines"]))
    totals = {"debit": float(sum(item["debit"] for item in entries)), "credit": float(sum(item["credit"] for item in entries))}
    balances: dict[str, Decimal] = {}
    for item in entries:
        for row in item["lines"]:
            balances[row["account"]] = balances.get(row["account"], Decimal("0")) + row["debit"] - row["credit"]
    checks = [
        {"label": "Cada asiento cuadra (debe = haber)", "ok": not unbalanced, "detail": f"Descuadrados: {unbalanced[:10]}" if unbalanced else None},
        {"label": "Sumas del diario iguales", "ok": abs(totals["debit"] - totals["credit"]) < 0.01, "detail": None},
        {"label": "Facturas con cuenta de gasto específica", "ok": default_account == 0 or None,
         "detail": f"{default_account} en {DEFAULT_ACCOUNT} (otros gastos): revisa la categoría" if default_account else None},
        {"label": "Terceros con NIF", "ok": no_tax_id == 0 or None, "detail": f"{no_tax_id} asiento(s) con proveedor o cliente sin NIF" if no_tax_id else None},
        {"label": "Todo lo del periodo contabilizado", "ok": not (pending or not_booked) or None,
         "detail": f"{len(not_booked)} factura(s) sin aprobar y {len(pending)} movimiento(s) sin conciliar" if pending or not_booked else None},
    ]
    return {
        "from": date_from.isoformat(), "to": date_to.isoformat(), "entries": entries, "count": len(entries), "totals": totals,
        "balances": [{"account": key, "name": ACCOUNTS.get(key, ""), "balance": float(value)} for key, value in sorted(balances.items())],
        "checks": checks, "pending": not_booked + pending,
        "status": "BORRADOR",
        "note": ("Borrador contable: revísalo con tu gestoría antes de darlo por definitivo. Casos aún por validar con datos reales: "
                 "rectificativas, anticipos, pagos parciales, impuestos y movimientos sin conciliar. Formato genérico con el Plan General "
                 "Contable; para A3, Sage u Holded con su formato propio hace falta su especificación de importación."),
    }


def rows(report: dict[str, Any]) -> list[list[Any]]:
    result = []
    for item in report["entries"]:
        for row in item["lines"]:
            result.append([item["number"], item["date"].isoformat(), row["account"], row["account_name"], row["tax_id"] or "", row["party"] or "",
                           item["concept"], f"{row['debit']:.2f}".replace(".", ","), f"{row['credit']:.2f}".replace(".", ","), item["document"] or ""])
    return result


def to_csv(report: dict[str, Any]) -> bytes:
    import csv
    import io

    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(COLUMNS)
    writer.writerows(rows(report))
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(report: dict[str, Any]) -> bytes:
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    sheet = book.active
    sheet.title = "Borrador libro diario"
    sheet.append([column.replace("_", " ").capitalize() for column in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for item in report["entries"]:
        for row in item["lines"]:
            sheet.append([item["number"], item["date"], row["account"], row["account_name"], row["tax_id"] or "", row["party"] or "",
                          item["concept"], float(row["debit"]), float(row["credit"]), item["document"] or ""])
    for column, width in zip("ABCDEFGHIJ", (8, 12, 8, 34, 12, 30, 50, 12, 12, 16)):
        sheet.column_dimensions[column].width = width
    for cells in sheet.iter_rows(min_row=2, min_col=8, max_col=9):
        for cell in cells:
            cell.number_format = "#,##0.00"
    balances = book.create_sheet("Saldos")
    balances.append(["Cuenta", "Nombre", "Saldo"])
    for item in report["balances"]:
        balances.append([item["account"], item["name"], item["balance"]])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def serialize(report: dict[str, Any], limit: int = 200) -> dict[str, Any]:
    def plain(item: dict[str, Any]) -> dict[str, Any]:
        return {**item, "date": item["date"].isoformat(), "debit": float(item["debit"]), "credit": float(item["credit"]),
                "lines": [{**row, "debit": float(row["debit"]), "credit": float(row["credit"])} for row in item["lines"]]}

    return {**report, "entries": [plain(item) for item in report["entries"][:limit]], "truncated": max(0, len(report["entries"]) - limit)}
