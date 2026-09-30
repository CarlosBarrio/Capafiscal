from __future__ import annotations

import io
from datetime import date
from datetime import timedelta

from openpyxl import load_workbook
from sqlalchemy import text

from app.database import add_missing_columns
from app.database import engine
from tests.conftest import FIXTURES_DIR
from tests.conftest import upload


def invoice_id(payload: dict) -> int:
    return payload["document"]["invoice"]["id"]


def test_upload_processes_invoice_and_detects_duplicate_file(client, sample_pdfs):
    first = upload(client, sample_pdfs["factura_repsol_2026"])

    assert first["duplicate"] is False
    assert first["document"]["status"] == "READY_FOR_APPROVAL"
    assert first["document"]["invoice"]["supplier_name"] == "REPSOL COMERCIAL S.A."

    second = upload(client, sample_pdfs["factura_repsol_2026"])

    assert second["duplicate"] is True
    assert len(client.get("/api/documents").json()) == 1


def test_mismatch_cannot_be_approved_until_corrected(client, sample_pdfs):
    payload = upload(client, sample_pdfs["factura_iberdrola_MAL"])
    iid = invoice_id(payload)

    assert payload["document"]["invoice"]["validation_status"] == "MISMATCH"

    response = client.post(f"/api/invoices/{iid}/approve", json={})
    assert response.status_code == 409

    response = client.patch(f"/api/invoices/{iid}", json={"total": "121.00"})
    assert response.status_code == 200
    assert response.json()["validation_status"] == "VALID"

    response = client.post(f"/api/invoices/{iid}/approve", json={})
    assert response.status_code == 200
    assert response.json()["invoice"]["review_status"] == "APPROVED"


def test_force_approval_with_missing_fields(client, sample_pdfs):
    iid = invoice_id(upload(client, sample_pdfs["factura_endesa_2026"]))
    client.patch(f"/api/invoices/{iid}", json={"supplier_name": None})

    response = client.post(f"/api/invoices/{iid}/approve", json={})
    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "MISSING_FIELDS"

    # Antes fallaba con NameError por un import que faltaba.
    response = client.post(f"/api/invoices/{iid}/approve", json={"force": True})
    assert response.status_code == 200
    assert response.json()["approved_with_warnings"] is True


def test_strong_duplicate_blocks_even_forced_approval(client, sample_pdfs, tmp_path):
    upload(client, sample_pdfs["factura_vodafone_2026"])

    copy = tmp_path / "vodafone_reenviada.txt"
    copy.write_text(
        "VODAFONE ESPANA S.A.U.\nCIF: A80907397\n"
        "Factura nº VF-2026-556231 Fecha 08/07/2026\n"
        "Base imponible 71,90\nIVA 21% 15,10\nTotal 87,00 EUR\n",
        encoding="utf-8",
    )
    duplicate = upload(client, copy)
    iid = invoice_id(duplicate)

    assert duplicate["document"]["invoice"]["duplicate_status"] == "STRONG"

    response = client.post(f"/api/invoices/{iid}/approve", json={"force": True})
    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "STRONG_DUPLICATE"


def test_reopen_and_reject_flow(client, sample_pdfs):
    iid = invoice_id(upload(client, sample_pdfs["factura_makro_2026"]))
    client.post(f"/api/invoices/{iid}/approve", json={})

    assert client.post(
        f"/api/invoices/{iid}/reject", json={"reason": "Error"}
    ).status_code == 409

    response = client.post(
        f"/api/invoices/{iid}/reopen", json={"reason": "Revisar precio"}
    )
    assert response.status_code == 200
    assert response.json()["invoice"]["review_status"] == "PENDING"

    assert client.post(
        f"/api/invoices/{iid}/reopen", json={"reason": "Otra vez"}
    ).status_code == 409

    response = client.post(
        f"/api/invoices/{iid}/reject", json={"reason": "Precio incorrecto"}
    )
    assert response.status_code == 200

    actions = [
        event["action"]
        for event in client.get(
            f"/api/documents/{response.json()['document']['id']}/audit"
        ).json()
    ]
    assert "invoice.reopened" in actions
    assert "invoice.rejected" in actions


def test_payments_flow_and_overview(client, sample_pdfs):
    iid = invoice_id(upload(client, sample_pdfs["factura_repsol_2026"]))

    response = client.post(f"/api/invoices/{iid}/payment", json={"paid": True})
    assert response.status_code == 409  # todavía no aprobada

    client.post(f"/api/invoices/{iid}/approve", json={})
    overdue_date = (date.today() - timedelta(days=5)).isoformat()
    client.post(f"/api/invoices/{iid}/reopen", json={"reason": "Añadir vencimiento"})
    client.patch(f"/api/invoices/{iid}", json={"due_date": overdue_date})
    client.post(f"/api/invoices/{iid}/approve", json={})

    overview = client.get("/api/payments").json()
    assert overview["unpaid_count"] == 1
    assert overview["overdue_count"] == 1
    assert overview["items"][0]["state"] == "OVERDUE"

    unpaid = client.get("/api/documents", params={"payment": "overdue"}).json()
    assert len(unpaid) == 1

    response = client.post(
        f"/api/invoices/{iid}/payment",
        json={"paid": True, "payment_method": "transferencia"},
    )
    assert response.status_code == 200
    assert response.json()["invoice"]["payment_method"] == "TRANSFERENCIA"
    assert client.get("/api/payments").json()["unpaid_count"] == 0

    response = client.post(
        f"/api/invoices/{iid}/payment",
        json={"paid": True, "payment_method": "bitcoin"},
    )
    assert response.status_code == 422


def approve_samples(client, sample_pdfs, stems):
    for stem in stems:
        iid = invoice_id(upload(client, sample_pdfs[stem]))
        assert client.post(f"/api/invoices/{iid}/approve", json={}).status_code == 200


def test_vat_report_and_ledger_export(client, sample_pdfs):
    approve_samples(
        client,
        sample_pdfs,
        ["factura_repsol_2026", "factura_endesa_2026", "factura_vodafone_2026"],
    )
    upload(client, sample_pdfs["factura_iberdrola_MAL"])  # queda pendiente

    report = client.get("/api/reports/vat", params={"year": 2026, "quarter": 3}).json()

    assert report["approved_invoices"] == 3
    assert report["pending_invoices"] == 1
    assert report["total_base"] == 726.78
    assert report["total_tax"] == 152.63
    assert report["by_rate"][0]["label"] == "21 %"
    assert report["warnings"]

    empty = client.get("/api/reports/vat", params={"year": 2026, "quarter": 1}).json()
    assert empty["approved_invoices"] == 0

    csv_response = client.get(
        "/api/exports/ledger", params={"format": "csv", "year": 2026, "quarter": 3}
    )
    assert csv_response.status_code == 200
    lines = csv_response.content.decode("utf-8-sig").strip().splitlines()
    assert lines[0].startswith("Nº orden;Fecha expedición")
    assert len(lines) == 4
    assert ";21;" in lines[1]

    xlsx_response = client.get(
        "/api/exports/ledger", params={"year": 2026, "quarter": 3}
    )
    workbook = load_workbook(io.BytesIO(xlsx_response.content))
    sheet = workbook.active
    assert sheet["A3"].value == "Nº orden"
    assert sheet.max_row == 7  # título, vacía, cabecera, 3 facturas, totales

    activity = client.get("/api/activity", params={"action": "ledger.*"}).json()
    assert len(activity) == 2


def test_search_filters_and_suppliers(client, sample_pdfs):
    approve_samples(client, sample_pdfs, ["factura_repsol_2026"])
    upload(client, sample_pdfs["factura_vodafone_2026"])

    assert len(client.get("/api/documents", params={"q": "vodafone"}).json()) == 1
    assert len(client.get("/api/documents", params={"q": "A28047223"}).json()) == 1
    assert len(
        client.get("/api/documents", params={"review_status": "APPROVED"}).json()
    ) == 1
    assert len(
        client.get("/api/documents", params={"category": "Combustible"}).json()
    ) == 1
    assert len(
        client.get(
            "/api/documents",
            params={"date_from": "2026-07-10", "date_to": "2026-07-31"},
        ).json()
    ) == 1

    suppliers = client.get("/api/suppliers").json()
    assert suppliers[0]["name"] == "REPSOL COMERCIAL S.A."
    assert suppliers[0]["total_approved"] == 479.5
    assert suppliers[0]["tax_id_valid"] is True
    assert len(client.get("/api/suppliers", params={"q": "voda"}).json()) == 1


def test_assistant_intents(client, sample_pdfs):
    approve_samples(client, sample_pdfs, ["factura_repsol_2026"])

    def ask(question: str) -> str:
        response = client.post("/api/assistant/query", json={"question": question})
        assert response.status_code == 200
        return response.json()["answer"]

    assert "83,22 €" in ask("IVA del 3T 2026")
    assert "sin pagar" in ask("¿Qué pagos tengo pendientes?")
    assert "REPSOL" in ask("cuánto he gastado en repsol")


def test_review_inbox_orders_by_real_priority(client, sample_pdfs):
    upload(client, sample_pdfs["factura_repsol_2026"])  # LOW
    upload(client, sample_pdfs["factura_iberdrola_MAL"])  # HIGH
    upload(client, FIXTURES_DIR / "bloques.txt")  # NORMAL o LOW

    priorities = [
        task["priority"]
        for task in client.get("/api/tasks/review-inbox").json()
    ]
    order = {"HIGH": 0, "NORMAL": 1, "LOW": 2}

    assert priorities[0] == "HIGH"
    assert priorities == sorted(priorities, key=order.get)


def test_missing_columns_are_added_to_existing_database(client):
    with engine.begin() as connection:
        connection.execute(text("ALTER TABLE invoices DROP COLUMN payment_method"))

    added = add_missing_columns()

    assert "invoices.payment_method" in added
    assert add_missing_columns() == []


def test_reprocess_pending_skips_approved(client, sample_pdfs):
    approved = invoice_id(upload(client, sample_pdfs["factura_repsol_2026"]))
    client.post(f"/api/invoices/{approved}/approve", json={})
    upload(client, sample_pdfs["factura_endesa_2026"])
    upload(client, sample_pdfs["requerimiento_aeat_2026"])

    result = client.post("/api/documents/reprocess-pending").json()

    assert result["success"] is True
    assert result["processed"] == 2

    statuses = {
        document["original_filename"]: document["status"]
        for document in client.get("/api/documents").json()
    }
    assert statuses["factura_repsol_2026.pdf"] == "APPROVED"
